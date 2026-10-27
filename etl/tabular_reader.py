"""Read supported CSV/XLSX tables without writing uploads to disk."""

from __future__ import annotations

from pathlib import Path
from typing import IO
from zipfile import ZipFile

import pandas as pd

from etl.models import DatasetValidationError

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_WORKBOOK_BYTES = 100 * 1024 * 1024
MAX_EXCEL_ROWS = 100_000
MAX_EXCEL_COLUMNS = 256
TableSource = str | Path | IO


def _rewind(source: TableSource) -> None:
    if hasattr(source, "seek"):
        source.seek(0)


def _open_workbook(source: TableSource):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise DatasetValidationError(
            "Excel support requires openpyxl. Run: python -m pip install -r requirements.txt"
        ) from error
    try:
        _rewind(source)
        with ZipFile(source) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_WORKBOOK_BYTES:
                raise DatasetValidationError("Excel workbook exceeds the 100 MB expanded-size limit.")
        _rewind(source)
        return load_workbook(source, read_only=True, data_only=False, keep_links=False)
    except DatasetValidationError:
        raise
    except Exception as error:
        raise DatasetValidationError(
            "Cannot read this Excel workbook. Upload an unencrypted, valid .xlsx file."
        ) from error


def workbook_sheets(source: TableSource) -> list[str]:
    """List sheet names; the caller must explicitly select a sheet for import."""
    workbook = _open_workbook(source)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()
        _rewind(source)


def read_security_table(
    source: TableSource, *, filename: str | None = None,
    sheet_name: str | None = None, nrows: int | None = None,
) -> pd.DataFrame:
    """Read one table. Excel row 1 is the header; formulas are not accepted."""
    name = filename or (str(source) if isinstance(source, (str, Path)) else getattr(source, "name", ""))
    suffix = Path(name).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise DatasetValidationError("Supported upload formats are .csv and .xlsx; save older .xls files as .xlsx.")
    _rewind(source)
    if suffix == ".csv":
        try:
            return pd.read_csv(source, nrows=nrows)
        except Exception as error:
            raise DatasetValidationError("Cannot read the CSV. Check its encoding and column structure.") from error
        finally:
            _rewind(source)

    workbook = _open_workbook(source)
    try:
        if sheet_name is None:
            if len(workbook.sheetnames) != 1:
                raise DatasetValidationError("This workbook has multiple worksheets. Select one worksheet to import.")
            sheet_name = workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            raise DatasetValidationError("The selected worksheet does not exist in this workbook.")
        sheet = workbook[sheet_name]
        if sheet.max_row and sheet.max_row > MAX_EXCEL_ROWS + 1:
            raise DatasetValidationError("Excel worksheets are limited to 100,000 data rows.")
        if sheet.max_column and sheet.max_column > MAX_EXCEL_COLUMNS:
            raise DatasetValidationError("Excel worksheets are limited to 256 columns.")
        rows = sheet.iter_rows()
        header = next(rows, ())
        if any(cell.data_type == "f" for cell in header):
            raise DatasetValidationError("Excel formulas are not supported. Replace formulas with values before uploading.")
        names = [cell.value for cell in header]
        while names and names[-1] is None:
            names.pop()
        if not names:
            raise DatasetValidationError("The selected worksheet is empty.")
        if any(not isinstance(name, str) or not name.strip() for name in names):
            raise DatasetValidationError("Excel row 1 must contain nonempty text column names.")
        if len(names) != len(set(names)):
            raise DatasetValidationError("Excel column names must be unique.")
        values = []
        for index, row in enumerate(rows):
            if nrows is not None and len(values) >= nrows:
                break
            if index >= MAX_EXCEL_ROWS or len(row) > MAX_EXCEL_COLUMNS:
                raise DatasetValidationError("The worksheet exceeds the supported row or column limit.")
            if any(cell.data_type == "f" for cell in row):
                raise DatasetValidationError("Excel formulas are not supported. Replace formulas with values before uploading.")
            data = [cell.value for cell in row]
            if any(value is not None for value in data[len(names):]):
                raise DatasetValidationError("Excel contains values in columns without a header.")
            data = data[:len(names)]
            if all(value is None for value in data):
                continue
            values.append(data)
        return pd.DataFrame(values, columns=names)
    except DatasetValidationError:
        raise
    except Exception as error:
        raise DatasetValidationError("Cannot parse the selected Excel worksheet.") from error
    finally:
        workbook.close()
        _rewind(source)
